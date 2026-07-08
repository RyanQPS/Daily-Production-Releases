#!/usr/bin/env python3
"""Rebuild Production_Release_Log.xlsx from release_state.json.

Zero formula errors — plain values only. Age-in-days is a simple integer computed
in Python (not =TODAY()-...), so the workbook renders identically wherever it's
opened.
"""
from __future__ import annotations

import datetime as dt
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
STATE = ROOT / "release_state.json"
OUT = ROOT / "Production_Release_Log.xlsx"

HEADERS = [
    "Job #", "Project Name", "Contractor", "PM", "Contact", "Phone",
    "Storm", "San", "Received", "Age (days)", "Notes", "Message-ID",
]

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
NEW_FILL = PatternFill("solid", fgColor="FFF2CC")  # new-since-last-run
OLD_FILL = PatternFill("solid", fgColor="F8CBAD")  # oldest pending


def load_state() -> dict:
    with STATE.open() as fh:
        return json.load(fh)


def parse_date(s: str | None) -> dt.date | None:
    if not s:
        return None
    return dt.date.fromisoformat(s[:10])


def build(today: dt.date, prior_ids: set[str]) -> pathlib.Path:
    state = load_state()
    pending = list(state.get("pending", []))
    ambiguous = list(state.get("ambiguous", []))

    # sort oldest -> newest so oldest job floats to the top row
    pending.sort(key=lambda r: r.get("received") or "9999-12-31")
    oldest_id = pending[0]["messageId"] if pending else None

    wb = Workbook()
    ws = wb.active
    ws.title = "Pending RTPs"

    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(pending, start=2):
        recv = parse_date(r.get("received"))
        age = (today - recv).days if recv else ""
        values = [
            r.get("jobNumber") or "",
            r.get("name") or "",
            r.get("contractor") or "",
            r.get("pm") or "",
            r.get("contact") or "",
            r.get("phone") or "",
            r.get("storm") or "",
            r.get("san") or "",
            r.get("received") or "",
            age,
            r.get("notes") or "",
            r.get("messageId") or "",
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        # highlight
        mid = r.get("messageId")
        if mid and mid == oldest_id:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=i, column=col).fill = OLD_FILL
        elif mid and mid not in prior_ids:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=i, column=col).fill = NEW_FILL

    widths = [10, 40, 24, 20, 30, 18, 8, 8, 12, 10, 60, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # Ambiguous sheet
    if ambiguous:
        aw = wb.create_sheet("Ambiguous")
        hdrs = ["Job #", "Project Name", "PM", "Received", "Reason", "Message-ID"]
        for col, h in enumerate(hdrs, start=1):
            c = aw.cell(row=1, column=col, value=h)
            c.fill = HDR_FILL
            c.font = HDR_FONT
        for i, r in enumerate(ambiguous, start=2):
            vals = [
                r.get("jobNumber") or "",
                r.get("name") or "",
                r.get("pm") or "",
                r.get("received") or "",
                r.get("reason") or "",
                r.get("messageId") or "",
            ]
            for col, v in enumerate(vals, start=1):
                c = aw.cell(row=i, column=col, value=v)
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for i, w in enumerate([12, 40, 20, 12, 60, 60], start=1):
            aw.column_dimensions[get_column_letter(i)].width = w
        aw.freeze_panes = "A2"

    # Summary sheet
    sw = wb.create_sheet("Summary")
    sw["A1"] = "Production Release Log — Summary"
    sw["A1"].font = Font(bold=True, size=14)
    sw["A3"] = "Generated"
    sw["B3"] = today.isoformat()
    sw["A4"] = "Total pending"
    sw["B4"] = len(pending)
    sw["A5"] = "Ambiguous / flagged"
    sw["B5"] = len(ambiguous)
    sw["A6"] = "Oldest pending (received)"
    sw["B6"] = pending[0]["received"] if pending else ""
    sw["A7"] = "Oldest pending (job)"
    sw["B7"] = f"{pending[0]['jobNumber']} — {pending[0]['name']}" if pending else ""

    # PM counts
    counts: dict[str, int] = {}
    for r in pending:
        counts[r.get("pm") or ""] = counts.get(r.get("pm") or "", 0) + 1
    sw["A9"] = "By PM"
    sw["A9"].font = Font(bold=True)
    row = 10
    for pm, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
        sw.cell(row=row, column=1, value=pm)
        sw.cell(row=row, column=2, value=n)
        row += 1
    for i, w in enumerate([30, 20], start=1):
        sw.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT)
    return OUT


def main() -> int:
    today = dt.date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else dt.date.today()
    prior_ids: set[str] = set()
    prior_path = ROOT / ".prior_ids.txt"
    if prior_path.exists():
        prior_ids = {ln.strip() for ln in prior_path.read_text().splitlines() if ln.strip()}
    out = build(today, prior_ids)
    print(f"Wrote {out}")
    # store current IDs for next run's diff
    state = load_state()
    ids = [r["messageId"] for r in state.get("pending", []) if r.get("messageId")]
    prior_path.write_text("\n".join(ids) + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
