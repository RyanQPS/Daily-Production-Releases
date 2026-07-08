#!/usr/bin/env python3
"""Force a recalc of Production_Release_Log.xlsx.

There are no formulas in the workbook (values are baked at build time), so
recalc is essentially a smoke test: open, re-save, verify no #REF!/#N/A/#DIV/0!
cells slipped in. Exits non-zero on any formula error string.
"""
from __future__ import annotations

import pathlib
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
BOOK = ROOT / "Production_Release_Log.xlsx"
ERROR_STRINGS = {"#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!"}


def main() -> int:
    if not BOOK.exists():
        print(f"Missing {BOOK}", file=sys.stderr)
        return 2
    wb = load_workbook(BOOK)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=False):
            for c in row:
                v = c.value
                if isinstance(v, str) and v in ERROR_STRINGS:
                    bad.append(f"{ws.title}!{c.coordinate} = {v}")
    if bad:
        print("Formula errors found:", *bad, sep="\n", file=sys.stderr)
        return 1
    wb.save(BOOK)
    print(f"Recalc OK — {BOOK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
